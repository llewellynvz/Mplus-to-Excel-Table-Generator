#!/usr/bin/env python
# -*- coding: utf-8 -*-

r"""
Creates a professional, APA-style Excel workbook FROM SCRATCH by parsing Mplus .out files.

Core behavior retained:
- Parses fit indices from Mplus "fit-card" sections (Chi-square, RMSEA, CFI/TLI, SRMR, Information Criteria).
- Parses STDYX standardized loadings and R-square values from Mplus output.
- Builds Table 1 (competing measurement models) and Table 2 (standardized factor loadings).
- Uses item labels exactly as they appear in the Mplus output.

New (formatting + usability only):
- Global formatting: Times New Roman 10, neat APA lines, consistent numeric formats.
- Adds an "Instructions" tab FIRST with factor-label mappings you can edit (updates factor labels in Table 2 via formulas).
- Table 2 now selects the TOP N models from Table 1 (default N=3) and builds one loading table per model.
- Boldfaces significant standardized loadings (λ only; SE not bold).
- Adds factor-level indices per model (AVE, alpha, omega; and omegaH/omegaS for bifactor-type models) computed from STDYX loadings,
  factor correlations (STDYX WITH section), and item residual variances derived from R2 (theta = 1 - R2).

Notes on reliability computations:
- Omega/alpha are computed for unit-weighted subscale scores, using the model-implied covariance matrix Σ = Λ Φ Λ' + Θ,
  where Λ are STDYX loadings, Φ are factor correlations from standardized output, and Θ are residual variances (diag(1 - R2)).
- Bifactor omegaH and omegaS are computed as the proportion of score variance attributable to the general (or specific) factor only,
  using the corresponding single-factor contribution to Σ. This matches standard "omega hierarchical" logic used in psychometric practice.

Usage:
  mplus2excel.py
  mplus2excel.py --input "C:\path\outs" --output "Tabulated Results_v7_6.xlsx" --top 3
"""

from __future__ import annotations

import argparse
import csv
import re
import logging
import sys
import traceback
from datetime import datetime
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Set

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter


# -----------------------------
# Logging & diagnostics
# -----------------------------
def setup_logging(folder: Path) -> Path:
    """Configure console + file logging.

    This does not alter analytic functionality; it only improves diagnostics.
    A 'run.log' file is created in the input folder.
    """
    log_path = folder / "run.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=[
            logging.FileHandler(log_path, encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )
    logging.info("Mplus2Excel run started")
    return log_path

# -----------------------------
# User-facing thresholds
# -----------------------------
P_BOLD_CUT = 0.05
LOADING_MIN_FOR_MEETS = 0.35
THETA_MAX_FOR_MEETS = 0.60  # uniqueness threshold (theta = residual variance); adjust if you prefer 0.80

# Fit thresholds (same spirit as earlier)
CFI_TLI_STRICT = 0.90
CFI_TLI_LOOSE  = 0.85
RMSEA_MAX = 0.08
SRMR_MAX  = 0.08
REQUIRE_PCLOSE_FOR_YES = True
PCLOSE_CUT = 0.05

# -----------------------------
# Parsing helpers (robust numeric)
# -----------------------------
NUM = r"[0-9.\-+EDed]+"
NUM_CLEAN = re.compile(r"^" + NUM + r"\*?$")  # allow trailing *

def _read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return path.read_text(encoding="latin-1", errors="ignore")

def _tofloat(x) -> Optional[float]:
    if x is None:
        return None
    try:
        s = str(x).strip().replace("D", "E").replace("d", "e").rstrip("*")
        return float(s)
    except Exception:
        return None

def _toint(x) -> Optional[int]:
    v = _tofloat(x)
    if v is None:
        return None
    try:
        return int(v)
    except Exception:
        return None

def _extract_after_heading(text: str, heading: str, window: int = 800) -> Optional[str]:
    idx = text.find(heading)
    if idx == -1:
        return None
    return text[idx: idx + window]

def _parse_model_id_and_type(filename: str) -> Tuple[Optional[int], str]:
    """Best-effort labeling when no manifest is provided."""
    base = filename
    model_id = None
    m = re.search(r"\bmm\s*([0-9]+)", base, re.I)
    if not m:
        m = re.search(r"(\d+)", base)
    if m:
        try:
            model_id = int(m.group(1))
        except Exception:
            model_id = None

    typ = base.rsplit(".out", 1)[0]
    typ = re.sub(r"^\s*mm\s*\d+\s*[-–]\s*", "", typ, flags=re.I)
    typ = re.sub(r"^\s*measurement\s*model\s*\d+\s*[-–]\s*", "", typ, flags=re.I)
    typ = typ.strip()
    return model_id, typ

def _natural_key(s: str):
    nums = re.findall(r"\d+", s)
    return (int(nums[0]) if nums else 10**9, s.lower())

# -----------------------------
# Data structures
# -----------------------------
@dataclass
class FitStats:
    chi2: Optional[float] = None
    df: Optional[int] = None
    p: Optional[float] = None
    cfi: Optional[float] = None
    tli: Optional[float] = None
    rmsea: Optional[float] = None
    rmsea_ci_lo: Optional[float] = None
    rmsea_ci_hi: Optional[float] = None
    pclose: Optional[float] = None
    srmr: Optional[float] = None
    aic: Optional[float] = None
    bic: Optional[float] = None
    abic: Optional[float] = None
    terminated: bool = True

@dataclass
class Loading:
    est: Optional[float] = None
    se: Optional[float] = None
    p: Optional[float] = None

@dataclass
class ParsedModel:
    out_file: Path
    model_label: str
    model_type: str
    fit: FitStats
    # observed item loadings: factor -> item -> Loading
    loadings: Dict[str, Dict[str, Loading]]
    # higher-order loadings: higher_factor -> first_order_factor -> Loading
    higher_order: Dict[str, Dict[str, Loading]]
    # factor correlations: (f1, f2) -> corr
    phi: Dict[Tuple[str, str], float]
    # item R2 and derived theta
    r2: Dict[str, float]
    role: str = ""
    order: int = 10**9
    model_id: Optional[int] = None

# -----------------------------
# Fit parsing
# -----------------------------
def parse_fit(text: str) -> FitStats:
    fs = FitStats()
    fs.terminated = ("THE MODEL ESTIMATION TERMINATED NORMALLY" in text)

    chi_block = _extract_after_heading(text, "Chi-Square Test of Model Fit", 600)
    if chi_block:
        m = re.search(r"Value\s+(" + NUM + r")\*?", chi_block)
        if m: fs.chi2 = _tofloat(m.group(1))
        m = re.search(r"Degrees of Freedom\s+(\d+)", chi_block)
        if m: fs.df = _toint(m.group(1))
        m = re.search(r"P-Value\s+(" + NUM + r")", chi_block)
        if m: fs.p = _tofloat(m.group(1))

    rm_block = _extract_after_heading(text, "RMSEA (Root Mean Square Error Of Approximation)", 700)
    if rm_block:
        m = re.search(r"Estimate\s+(" + NUM + r")\*?", rm_block)
        if m: fs.rmsea = _tofloat(m.group(1))
        m = re.search(r"90 Percent C\.I\.\s+(" + NUM + r")\s+(" + NUM + r")", rm_block)
        if m:
            fs.rmsea_ci_lo = _tofloat(m.group(1))
            fs.rmsea_ci_hi = _tofloat(m.group(2))
        m = re.search(r"Probability RMSEA <= \.05\s+(" + NUM + r")", rm_block)
        if m: fs.pclose = _tofloat(m.group(1))

    cfi_block = _extract_after_heading(text, "CFI/TLI", 400)
    if cfi_block:
        m = re.search(r"\n\s*CFI\s+(" + NUM + r")\*?", cfi_block)
        if m: fs.cfi = _tofloat(m.group(1))
        m = re.search(r"\n\s*TLI\s+(" + NUM + r")\*?", cfi_block)
        if m: fs.tli = _tofloat(m.group(1))
    else:
        m = re.search(r"\bCFI\s+(" + NUM + r")\*?", text)
        if m: fs.cfi = _tofloat(m.group(1))
        m = re.search(r"\bTLI\s+(" + NUM + r")\*?", text)
        if m: fs.tli = _tofloat(m.group(1))

    srmr_block = _extract_after_heading(text, "Standardized Root Mean Square Residual", 300)
    if srmr_block:
        m = re.search(r"Value\s+(" + NUM + r")\*?", srmr_block)
        if m: fs.srmr = _tofloat(m.group(1))
    if fs.srmr is None:
        m = re.search(r"\bSRMR\s+(" + NUM + r")\*?", text)
        if m: fs.srmr = _tofloat(m.group(1))

    ic_block = _extract_after_heading(text, "Information Criteria", 900)
    if ic_block:
        m = re.search(r"Akaike \(AIC\)\s+(" + NUM + r")\*?", ic_block)
        if m: fs.aic = _tofloat(m.group(1))
        m = re.search(r"Bayesian \(BIC\)\s+(" + NUM + r")\*?", ic_block)
        if m: fs.bic = _tofloat(m.group(1))
        m = re.search(r"Sample-Size Adjusted BIC\s+(" + NUM + r")\*?", ic_block)
        if m: fs.abic = _tofloat(m.group(1))
    return fs

# -----------------------------
# Standardized section extraction
# -----------------------------
def _extract_after_header(text: str, header_pat: str) -> Optional[str]:
    m = re.search(header_pat, text, re.I)
    if not m:
        return None
    tail = text[m.start():]
    stops = [
        r"\nSTDY(?!X)",
        r"\nMODEL RESULTS",
        r"\nMODEL FIT INFORMATION",
        r"\nTECHNICAL",
        r"\nQUALITY OF NUMERICAL RESULTS",
        r"\nINPUT READING TERMINATED",
        r"\nEND OF OUTPUT",
    ]
    for sp in stops:
        m2 = re.search(sp, tail, re.I)
        if m2 and m2.start() > 20:
            return tail[:m2.start()]
    return tail

def parse_usevariables(text: str) -> List[str]:
    m = re.search(r"USEVARIABLES\s+ARE\s+(.*?);", text, re.S | re.I)
    if not m:
        return []
    block = m.group(1)
    return re.findall(r"[A-Za-z0-9_]+", block)

def parse_rsquare(text: str) -> Dict[str, float]:
    r2: Dict[str, float] = {}
    m = re.search(r"\nR-SQUARE\s*\n", text, re.I)
    if not m:
        return r2
    tail = text[m.end():]
    stop = re.search(r"\n\s*\n\s*\n|\nMODEL FIT INFORMATION|\nMODEL RESULTS|\nSTANDARDIZED MODEL RESULTS|\nTECHNICAL", tail, re.I)
    sec = tail if stop is None else tail[:stop.start()]
    for ln in sec.splitlines():
        parts = ln.strip().split()
        if len(parts) >= 2 and re.fullmatch(r"[A-Za-z0-9_]+", parts[0]) and NUM_CLEAN.match(parts[1] or ""):
            r2[parts[0].strip().strip(",")] = float(parts[1].rstrip("*"))
    return r2

def parse_stdyx_loadings(text: str) -> Dict[str, Dict[str, Loading]]:
    """Return factor -> indicator -> Loading from STDYX (includes latent indicators if present)."""
    loadings: Dict[str, Dict[str, Loading]] = {}
    sec = _extract_after_header(text, r"STDYX\s+Standardization")
    if sec is None:
        sec = _extract_after_header(text, r"STANDARDIZED MODEL RESULTS")
    if sec is None:
        return loadings

    lines = sec.splitlines()
    current_factor: Optional[str] = None
    factor_pat = re.compile(r"^\s*([A-Za-z][A-Za-z0-9_]*)\s+BY\s*$", re.I)

    for ln in lines:
        m = factor_pat.match(ln)
        if m:
            current_factor = m.group(1)
            loadings.setdefault(current_factor, {})
            continue

        if current_factor is None:
            continue

        parts = ln.strip().split()
        if len(parts) < 3:
            if ln.strip() == "":
                current_factor = None
            continue

        ind = parts[0]
        if not re.fullmatch(r"[A-Za-z0-9_]+", ind):
            continue

        nums = [p.rstrip("*") for p in parts[1:] if NUM_CLEAN.match(p)]
        if len(nums) >= 2:
            est = _tofloat(nums[0])
            se  = _tofloat(nums[1])
            p   = _tofloat(nums[3]) if len(nums) >= 4 else None
            loadings[current_factor][ind] = Loading(est=est, se=se, p=p)

    return {f: d for f, d in loadings.items() if d}

def parse_stdyx_factor_correlations(text: str) -> Dict[Tuple[str, str], float]:
    """Parse factor correlations from the STDYX section (WITH blocks)."""
    phi: Dict[Tuple[str, str], float] = {}
    sec = _extract_after_header(text, r"STDYX\s+Standardization")
    if sec is None:
        sec = _extract_after_header(text, r"STANDARDIZED MODEL RESULTS")
    if sec is None:
        return phi

    current: Optional[str] = None
    with_pat = re.compile(r"^\s*([A-Za-z][A-Za-z0-9_]*)\s+WITH\s*$", re.I)
    for ln in sec.splitlines():
        m = with_pat.match(ln)
        if m:
            current = m.group(1)
            continue
        if current is None:
            continue
        parts = ln.strip().split()
        if len(parts) < 2:
            if ln.strip() == "":
                current = None
            continue
        other = parts[0]
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", other):
            continue
        nums = [p.rstrip("*") for p in parts[1:] if NUM_CLEAN.match(p)]
        if len(nums) >= 1:
            est = _tofloat(nums[0])
            if est is None:
                continue
            a, b = current, other
            if a == b:
                continue
            # store symmetric
            phi[(a, b)] = est
            phi[(b, a)] = est
    return phi

# -----------------------------
# Model classification & selection
# -----------------------------
def classify_meets_criteria(fs: FitStats) -> str:
    cfi = fs.cfi
    tli = fs.tli
    rmsea = fs.rmsea
    hi = fs.rmsea_ci_hi
    srmr = fs.srmr
    pclose = fs.pclose

    if any(v is None for v in [cfi, tli, rmsea, hi, srmr]):
        return "No"

    strict = (
        cfi >= CFI_TLI_STRICT and
        tli >= CFI_TLI_STRICT and
        rmsea <= RMSEA_MAX and
        hi <= RMSEA_MAX and
        srmr <= SRMR_MAX
    )
    pclose_ok = (pclose is None) or (pclose >= PCLOSE_CUT)

    if strict and ((not REQUIRE_PCLOSE_FOR_YES) or pclose_ok):
        return "Yes"

    failures = 0
    if cfi < CFI_TLI_STRICT: failures += 1
    if tli < CFI_TLI_STRICT: failures += 1
    if rmsea > RMSEA_MAX: failures += 1
    if hi > RMSEA_MAX: failures += 1
    if srmr > SRMR_MAX: failures += 1
    if REQUIRE_PCLOSE_FOR_YES and (pclose is not None and pclose < PCLOSE_CUT): failures += 1

    if failures <= 2 and cfi >= CFI_TLI_LOOSE and tli >= CFI_TLI_LOOSE:
        return "Marginally"
    return "No"

def _meets_rank(s: str) -> int:
    return {"Yes": 2, "Marginally": 1, "No": 0}.get(s, 0)

def select_top_models(models: List[ParsedModel], top_n: int = 3) -> List[ParsedModel]:
    """Rank models using fit indices and return the top N that have parsed loadings."""
    def key(m: ParsedModel):
        fs = m.fit
        meets = classify_meets_criteria(fs)
        chi2df = None
        if fs.chi2 is not None and fs.df:
            chi2df = fs.chi2 / fs.df if fs.df != 0 else None
        return (
            -_meets_rank(meets),
            -(fs.cfi if fs.cfi is not None else -1.0),
            -(fs.tli if fs.tli is not None else -1.0),
            (fs.rmsea if fs.rmsea is not None else 9.0),
            (fs.srmr if fs.srmr is not None else 9.0),
            (chi2df if chi2df is not None else 9e9),
            (fs.abic if fs.abic is not None else 9e99),
            (fs.bic if fs.bic is not None else 9e99),
            (fs.aic if fs.aic is not None else 9e99),
            (m.order, m.out_file.name.lower()),
        )

    eligible = [m for m in models if m.fit.terminated and bool(m.loadings)]
    eligible.sort(key=key)
    return eligible[:max(1, top_n)]

def is_bifactor_like(model: ParsedModel) -> bool:
    s = ((model.model_type or "") + " " + model.out_file.stem).lower()
    # Avoid false positives for ESEM (cross-loadings make every factor load on every item).
    return ("bifactor" in s) or ("bi-factor" in s) or ("bif " in s) or (" bif" in s)

def infer_general_factor(model: ParsedModel) -> Optional[str]:
    if not model.loadings:
        return None
    # explicit names
    for fn in model.loadings.keys():
        up = fn.upper()
        if up in {"G", "GF", "GFACTOR", "GENERAL"} or "GENERAL" in up:
            return fn
    # count-based
    items = {it for d in model.loadings.values() for it in d.keys()}
    if not items:
        return None
    best = max(model.loadings.keys(), key=lambda k: len(model.loadings.get(k, {})))
    if len(model.loadings.get(best, {})) >= int(0.90 * len(items)):
        return best
    return None

def assign_primary_factor(model: ParsedModel, general_factor: Optional[str] = None) -> Dict[str, str]:
    """Assign each item to the factor with the largest absolute loading (optionally excluding general)."""
    best: Dict[str, Tuple[str, float]] = {}
    for f, items in model.loadings.items():
        if general_factor and f == general_factor:
            continue
        for it, ld in items.items():
            if ld.est is None:
                continue
            val = abs(ld.est)
            if it not in best or val > best[it][1]:
                best[it] = (f, val)
    return {it: f for it, (f, _) in best.items()}

# -----------------------------
# Reliability and AVE calculations from Λ, Φ, Θ
# -----------------------------
def _build_lambda_phi_theta(model: ParsedModel) -> Tuple[List[str], List[str], np.ndarray, np.ndarray, np.ndarray]:
    """Return items, factors, Lambda(p x q), Phi(q x q), Theta(p,)"""
    factors = list(model.loadings.keys())
    items = sorted({it for d in model.loadings.values() for it in d.keys()}, key=_natural_key)
    q = len(factors)
    p = len(items)
    idx_f = {f:i for i,f in enumerate(factors)}
    idx_i = {it:i for i,it in enumerate(items)}

    L = np.zeros((p, q), dtype=float)
    for f, d in model.loadings.items():
        j = idx_f[f]
        for it, ld in d.items():
            i = idx_i[it]
            if ld.est is not None:
                L[i, j] = float(ld.est)

    Phi = np.eye(q, dtype=float)
    for (a, b), v in model.phi.items():
        if a in idx_f and b in idx_f and a != b and v is not None:
            Phi[idx_f[a], idx_f[b]] = float(v)

    Theta = np.full((p,), np.nan, dtype=float)
    for it, r2 in model.r2.items():
        if it in idx_i and r2 is not None:
            Theta[idx_i[it]] = max(0.0, 1.0 - float(r2))
    return items, factors, L, Phi, Theta

def _model_implied_cov(L: np.ndarray, Phi: np.ndarray, Theta: np.ndarray) -> Optional[np.ndarray]:
    if np.any(np.isnan(Theta)):
        return None
    return (L @ Phi @ L.T) + np.diag(Theta)

def _alpha_from_cov(S: np.ndarray) -> Optional[float]:
    k = S.shape[0]
    if k < 2:
        return None
    tr = float(np.trace(S))
    tot = float(S.sum())
    if tot <= 0:
        return None
    return (k / (k - 1.0)) * (1.0 - (tr / tot))

def _omega_total_from_parts(S: np.ndarray, S_common: np.ndarray, w: np.ndarray) -> Optional[float]:
    tot = float(w.T @ S @ w)
    if tot <= 0:
        return None
    com = float(w.T @ S_common @ w)
    return com / tot

def _contribution_cov(L: np.ndarray, Phi: np.ndarray, cols: List[int]) -> np.ndarray:
    """Return covariance contribution from selected factor columns."""
    if len(cols) == 0:
        return np.zeros((L.shape[0], L.shape[0]), dtype=float)
    Lc = L[:, cols]
    Phic = Phi[np.ix_(cols, cols)]
    return Lc @ Phic @ Lc.T

def compute_factor_indices(model: ParsedModel) -> Dict[str, Dict[str, Optional[float]]]:
    """
    Compute indices for each primary factor (subscale items assigned by max |loading|):
      AVE (Fornell-Larcker): sum(lambda^2)/(sum(lambda^2)+sum(theta))
      Alpha: from model-implied cov of subscale items
      Omega: omega_total for unit-weighted subscale score
    For bifactor-like models, also compute:
      omegaH (general contribution to subscale score)
      omegaS (specific contribution to subscale score)
    """
    out: Dict[str, Dict[str, Optional[float]]] = {}

    items, factors, L, Phi, Theta = _build_lambda_phi_theta(model)
    S = _model_implied_cov(L, Phi, Theta)
    if S is None:
        # cannot compute without all thetas (R2)
        return out

    bif = is_bifactor_like(model)
    g_name = infer_general_factor(model) if bif else None
    primary = assign_primary_factor(model, general_factor=g_name if bif else None)

    idx_f = {f:i for i,f in enumerate(factors)}
    idx_i = {it:i for i,it in enumerate(items)}

    # common covariance for all factors
    S_common_all = (L @ Phi @ L.T)

    # second-order: compute a "higher-order" factor loading on first-order factors if present
    # We'll compute omega for that higher-order factor across ALL items (approx, via product of coefficients).
    if model.higher_order:
        for ho, firsts in model.higher_order.items():
            # implied item loadings for HO: lambda_item_ho = lambda_item_f * gamma_f
            implied = np.zeros((len(items),), dtype=float)
            for f1, ld_ho in firsts.items():
                if ld_ho.est is None:
                    continue
                if f1 not in idx_f:
                    continue
                j = idx_f[f1]
                gamma = float(ld_ho.est)
                implied += L[:, j] * gamma
            # treat as single-factor contribution
            S_ho_common = np.outer(implied, implied)  # var(ho)=1
            w_all = np.ones((len(items),), dtype=float)
            omega_ho = _omega_total_from_parts(S, S_ho_common, w_all)
            alpha_ho = _alpha_from_cov(S)
            # AVE for HO is not standard; omit
            out[f"HigherOrder:{ho}"] = {"k": float(len(items)), "AVE": None, "alpha": alpha_ho, "omega": omega_ho, "omegaH": None, "omegaS": None}

    # per primary factor
    for f in factors:
        if bif and g_name and f == g_name:
            continue
        its = [it for it, pf in primary.items() if pf == f]
        its = sorted(its, key=_natural_key)
        if len(its) == 0:
            continue

        rows = [idx_i[it] for it in its]
        cols_all = list(range(len(factors)))
        cols_spec = [idx_f[f]]
        cols_g = [idx_f[g_name]] if (bif and g_name and g_name in idx_f) else []

        S_sub = S[np.ix_(rows, rows)]
        S_common_sub = S_common_all[np.ix_(rows, rows)]
        w = np.ones((len(rows),), dtype=float)

        alpha = _alpha_from_cov(S_sub)
        omega = _omega_total_from_parts(S_sub, S_common_sub, w)

        # AVE from target loadings on the primary factor
        lam = L[rows, idx_f[f]]
        the = Theta[rows]
        if np.any(np.isnan(the)):
            ave = None
        else:
            num = float(np.sum(lam * lam))
            den = num + float(np.sum(the))
            ave = (num / den) if den > 0 else None

        omegaH = None
        omegaS = None
        if bif and g_name and cols_g:
            S_g = _contribution_cov(L[np.ix_(rows, cols_g)], np.eye(1), [0])  # since we passed only g column
            omegaH = _omega_total_from_parts(S_sub, S_g, w)
            S_s = _contribution_cov(L[np.ix_(rows, cols_spec)], np.eye(1), [0])
            omegaS = _omega_total_from_parts(S_sub, S_s, w)

        out[f] = {"k": float(len(its)), "AVE": ave, "alpha": alpha, "omega": omega, "omegaH": omegaH, "omegaS": omegaS}
    return out

# -----------------------------
# Manifest handling (optional)
# -----------------------------
def load_manifest(folder: Path) -> Dict[str, Dict[str, str]]:
    path = folder / "model_manifest.csv"
    if not path.exists():
        return {}
    out: Dict[str, Dict[str, str]] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if not row.get("out_file"):
                continue
            key = row["out_file"].strip().lower()
            out[key] = {k: (v.strip() if isinstance(v, str) else "") for k, v in row.items()}
    return out

# -----------------------------
# Excel style
# -----------------------------
FONT_BODY  = Font(name="Times New Roman", size=10)
FONT_BOLD  = Font(name="Times New Roman", size=10, bold=True)
FONT_TITLE = Font(name="Times New Roman", size=12, bold=True)
FONT_NOTE  = Font(name="Times New Roman", size=9)

ALIGN_LEFT   = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_RIGHT  = Alignment(horizontal="right", vertical="center", wrap_text=True)

LINE = Side(style="thin", color="000000")

def border_top(): return Border(top=LINE)
def border_bottom(): return Border(bottom=LINE)
def border_top_bottom(): return Border(top=LINE, bottom=LINE)

def _set(ws, r, c, val, font=None, align=None, border: Optional[Border]=None, num_format: Optional[str]=None):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font = font if font is not None else FONT_BODY
    cell.alignment = align if align is not None else ALIGN_CENTER
    if border is not None:
        cell.border = border
    if num_format is not None and val is not None:
        cell.number_format = num_format
    return cell

def _apply_border_row(ws, r, c1, c2, border: Border):
    for c in range(c1, c2 + 1):
        ws.cell(r, c).border = border

def _set_col_widths(ws, widths: Dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

# -----------------------------
# Workbook builders
# -----------------------------
def build_instructions(ws, factor_names: List[str]):
    ws.title = "Instructions"
    _set(ws, 1, 1, "Instructions", font=FONT_TITLE, align=ALIGN_LEFT)
    _set(ws, 3, 1, "Factor label mapping (edit Column B to rename factors across Table 2):", font=FONT_BOLD, align=ALIGN_LEFT)

    _set(ws, 5, 1, "Original", font=FONT_BOLD)
    _set(ws, 5, 2, "Display label", font=FONT_BOLD)
    _apply_border_row(ws, 5, 1, 2, border_top_bottom())

    r = 6
    for f in factor_names:
        _set(ws, r, 1, f, align=ALIGN_LEFT)
        _set(ws, r, 2, f, align=ALIGN_LEFT)
        r += 1
    _apply_border_row(ws, r - 1, 1, 2, border_bottom())

    # Thresholds
    trow = 3
    _set(ws, trow, 4, "Thresholds (used when the workbook is generated):", font=FONT_BOLD, align=ALIGN_LEFT)
    _set(ws, trow + 2, 4, "p-value for bold λ", font=FONT_BODY, align=ALIGN_LEFT)
    _set(ws, trow + 2, 5, P_BOLD_CUT, num_format="0.00", align=ALIGN_RIGHT)
    _set(ws, trow + 3, 4, "min |λ| for Meets", font=FONT_BODY, align=ALIGN_LEFT)
    _set(ws, trow + 3, 5, LOADING_MIN_FOR_MEETS, num_format="0.00", align=ALIGN_RIGHT)
    _set(ws, trow + 4, 4, "max θ (residual var) for Meets", font=FONT_BODY, align=ALIGN_LEFT)
    _set(ws, trow + 4, 5, THETA_MAX_FOR_MEETS, num_format="0.00", align=ALIGN_RIGHT)

    _set_col_widths(ws, {1: 20, 2: 30, 4: 28, 5: 10})
    ws.freeze_panes = "A6"

def build_table1(ws, models: List[ParsedModel]):
    ws.title = "Table 1. Measurement Models"
    ws.freeze_panes = "A5"

    _set(ws, 1, 1, "Table 1", font=FONT_TITLE, align=ALIGN_LEFT)
    _set(ws, 2, 1, "Competing Measurement Models", font=FONT_TITLE, align=ALIGN_LEFT)

    headers = ["Model", "Type", "χ²", "df", "p(χ²)", "CFI", "TLI", "RMSEA", "90% CI", "pclose", "SRMR", "AIC", "BIC", "aBIC", "Meets criteria", "Notes", "Filename"]
    for i, h in enumerate(headers, start=1):
        _set(ws, 4, i, h, font=FONT_BOLD, align=ALIGN_CENTER)

    _apply_border_row(ws, 4, 1, len(headers), border_top_bottom())

    widths = {1: 10, 2: 28, 3: 10, 4: 6, 5: 10, 6: 7, 7: 7, 8: 8, 9: 14, 10: 8, 11: 8,
              12: 12, 13: 12, 14: 12, 15: 14, 16: 26, 17: 40}
    _set_col_widths(ws, widths)

    r = 5
    for pm in models:
        fs = pm.fit
        meets = classify_meets_criteria(fs)
        notes = []
        if not fs.terminated:
            notes.append("Did not terminate normally")
        note = "; ".join(notes) if notes else ""

        ci = ""
        if fs.rmsea_ci_lo is not None and fs.rmsea_ci_hi is not None:
            ci = f"[{fs.rmsea_ci_lo:.3f}-{fs.rmsea_ci_hi:.3f}]"

        _set(ws, r, 1, pm.model_label, align=ALIGN_LEFT)
        _set(ws, r, 2, pm.model_type, align=ALIGN_LEFT)

        _set(ws, r, 3, fs.chi2, align=ALIGN_RIGHT, num_format="0.000")
        _set(ws, r, 4, fs.df, align=ALIGN_RIGHT, num_format="0")
        _set(ws, r, 5, fs.p, align=ALIGN_RIGHT, num_format="0.0000")
        _set(ws, r, 6, fs.cfi, align=ALIGN_RIGHT, num_format="0.000")
        _set(ws, r, 7, fs.tli, align=ALIGN_RIGHT, num_format="0.000")
        _set(ws, r, 8, fs.rmsea, align=ALIGN_RIGHT, num_format="0.000")
        _set(ws, r, 9, ci, align=ALIGN_CENTER)
        _set(ws, r, 10, fs.pclose, align=ALIGN_RIGHT, num_format="0.0000")
        _set(ws, r, 11, fs.srmr, align=ALIGN_RIGHT, num_format="0.000")
        _set(ws, r, 12, fs.aic, align=ALIGN_RIGHT, num_format="0.000")
        _set(ws, r, 13, fs.bic, align=ALIGN_RIGHT, num_format="0.000")
        _set(ws, r, 14, fs.abic, align=ALIGN_RIGHT, num_format="0.000")
        _set(ws, r, 15, meets, align=ALIGN_CENTER)
        _set(ws, r, 16, note, align=ALIGN_LEFT)
        _set(ws, r, 17, pm.out_file.name, align=ALIGN_LEFT)
        r += 1

    if r > 5:
        _apply_border_row(ws, r - 1, 1, len(headers), border_bottom())

    note_r = r + 1
    ws.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=len(headers))
    note = ("Note. χ² = chi-square; df = degrees of freedom; RMSEA = root mean square error of approximation; "
            "CI = confidence interval; SRMR = standardized root mean square residual; aBIC = sample-size adjusted BIC. "
            "Meets criteria is a heuristic summary and should be interpreted alongside theory and diagnostics.")
    _set(ws, note_r, 1, note, font=FONT_NOTE, align=ALIGN_LEFT)

def _factor_display_formula(factor: str, map_range: str = "Instructions!$A$6:$B$200") -> str:
    # Excel formula: lookup original in mapping table; fall back to original.
    return f'=IFERROR(VLOOKUP("{factor}",{map_range},2,FALSE),"{factor}")'

def build_table2(ws, models: List[ParsedModel], top_models: List[ParsedModel]):
    ws.title = "Table 2. Std Factor Loadings"
    ws.freeze_panes = "A6"

    _set(ws, 1, 1, "Table 2", font=FONT_TITLE, align=ALIGN_LEFT)
    _set(ws, 2, 1, "Standardized Factor Loadings (Top models)", font=FONT_TITLE, align=ALIGN_LEFT)

    r = 4
    for idx, model in enumerate(top_models, start=1):
        # Section heading
        title = f"Model {model.model_id if model.model_id is not None else idx}: {model.model_type}"
        _set(ws, r, 1, title, font=FONT_BOLD, align=ALIGN_LEFT)
        r += 1

        # Build factor set/order and item assignment
        bif = is_bifactor_like(model)
        g_name = infer_general_factor(model) if bif else None

        # factors in display order
        factors = list(model.loadings.keys())
        if bif and g_name and g_name in factors:
            factors = [g_name] + [f for f in factors if f != g_name]

        primary = assign_primary_factor(model, general_factor=g_name if bif else None)
        factor_groups = [f for f in factors if (not (bif and g_name and f == g_name)) and any(primary.get(it) == f for it in primary)]

        # all observed items
        items_all = sorted({it for d in model.loadings.values() for it in d.keys()}, key=_natural_key)
        items_by_factor = {f: sorted([it for it in items_all if primary.get(it) == f], key=_natural_key) for f in factor_groups}

        # Column layout:
        # A Factor | B Item | then (λ, SE) for each factor | R2 | θ | Meets | d
        col_factor = 1
        col_item = 2
        start = 3
        factor_cols = {}
        c = start
        for f in factors:
            factor_cols[f] = (c, c+1)  # (lambda, se)
            c += 2
        col_r2 = c
        col_theta = c + 1
        col_meets = c + 2
        col_d = c + 3
        last_col = col_d

        # Header row
        _set(ws, r, col_factor, "Factor", font=FONT_BOLD)
        _set(ws, r, col_item, "Item", font=FONT_BOLD)
        # merged factor headers (row r)
        for f in factors:
            lcol, scol = factor_cols[f]
            ws.merge_cells(start_row=r, start_column=lcol, end_row=r, end_column=scol)
            _set(ws, r, lcol, _factor_display_formula(f), font=FONT_BOLD)
        _set(ws, r, col_r2, "R²", font=FONT_BOLD)
        _set(ws, r, col_theta, "θ", font=FONT_BOLD)
        _set(ws, r, col_meets, "Meets", font=FONT_BOLD)
        _set(ws, r, col_d, "d", font=FONT_BOLD)

        _apply_border_row(ws, r, 1, last_col, border_top_bottom())
        r += 1

        # Subheader row: λ / SE per factor
        _set(ws, r, col_factor, "", font=FONT_BOLD)
        _set(ws, r, col_item, "", font=FONT_BOLD)
        for f in factors:
            lcol, scol = factor_cols[f]
            _set(ws, r, lcol, "λ", font=FONT_BOLD)
            _set(ws, r, scol, "SE", font=FONT_BOLD)
        _set(ws, r, col_r2, "", font=FONT_BOLD)
        _set(ws, r, col_theta, "", font=FONT_BOLD)
        _set(ws, r, col_meets, "", font=FONT_BOLD)
        _set(ws, r, col_d, "", font=FONT_BOLD)
        _apply_border_row(ws, r, 1, last_col, border_bottom())
        r += 1

        # Column widths (reasonable defaults)
        widths = {1: 20, 2: 14}
        for f in factors:
            lcol, scol = factor_cols[f]
            widths[lcol] = 7
            widths[scol] = 7
        widths[col_r2] = 7
        widths[col_theta] = 7
        widths[col_meets] = 9
        widths[col_d] = 5
        _set_col_widths(ws, widths)

        # Body: grouped by primary factor
        for fgrp in factor_groups:
            # factor header row (display label via formula)
            _set(ws, r, col_factor, _factor_display_formula(fgrp), font=FONT_BOLD, align=ALIGN_LEFT)
            for cc in range(col_item, last_col+1):
                _set(ws, r, cc, "", font=FONT_BOLD)
            r += 1

            for it in items_by_factor.get(fgrp, []):
                _set(ws, r, col_factor, "", align=ALIGN_LEFT)
                _set(ws, r, col_item, it, align=ALIGN_LEFT)

                # fill loadings for ALL factors (ESEM-style matrix; CFA will have blanks)
                meets_primary = True
                # theta for meets
                r2 = model.r2.get(it)
                theta = (1.0 - r2) if (r2 is not None) else None

                for f in factors:
                    ld = model.loadings.get(f, {}).get(it)
                    lcol, scol = factor_cols[f]
                    if ld and ld.est is not None:
                        c_l = _set(ws, r, lcol, float(ld.est), align=ALIGN_RIGHT, num_format="0.000")
                        # significance -> bold lambda only
                        sig = False
                        if ld.p is not None:
                            sig = (ld.p < P_BOLD_CUT)
                        elif ld.se is not None and ld.se != 0:
                            z = abs(ld.est / ld.se)
                            sig = (z >= 1.96)
                        if sig:
                            c_l.font = Font(name="Times New Roman", size=10, bold=True)

                        if ld.se is not None:
                            _set(ws, r, scol, float(ld.se), align=ALIGN_RIGHT, num_format="0.000")
                    # else blank cells will be created later by borders if needed

                if r2 is not None:
                    _set(ws, r, col_r2, float(r2), align=ALIGN_RIGHT, num_format="0.000")
                if theta is not None:
                    _set(ws, r, col_theta, float(theta), align=ALIGN_RIGHT, num_format="0.000")

                # Meets based on PRIMARY loading on fgrp
                ld_primary = model.loadings.get(fgrp, {}).get(it)
                meets = "No"
                if ld_primary and ld_primary.est is not None:
                    sig_primary = False
                    if ld_primary.p is not None:
                        sig_primary = (ld_primary.p < P_BOLD_CUT)
                    elif ld_primary.se is not None and ld_primary.se != 0:
                        sig_primary = (abs(ld_primary.est / ld_primary.se) >= 1.96)

                    cond1 = sig_primary
                    cond2 = (abs(ld_primary.est) >= LOADING_MIN_FOR_MEETS)
                    cond3 = (theta is None) or (theta <= THETA_MAX_FOR_MEETS)
                    if cond1 and cond2 and cond3:
                        meets = "Yes"
                    elif cond1 and cond2:
                        meets = "Marg."
                _set(ws, r, col_meets, meets, align=ALIGN_CENTER)
                _set(ws, r, col_d, "", align=ALIGN_CENTER)

                r += 1

        # bottom border under the model section body
        _apply_border_row(ws, r - 1, 1, last_col, border_bottom())

        # Factor-level indices block
        r += 1
        _set(ws, r, 1, "Factor-level indices (unit-weighted scores)", font=FONT_BOLD, align=ALIGN_LEFT)
        r += 1
        headers = ["Factor", "k", "AVE", "α", "ω", "ωH", "ωS"]
        for j, h in enumerate(headers, start=1):
            _set(ws, r, j, h, font=FONT_BOLD)
        _apply_border_row(ws, r, 1, len(headers), border_top_bottom())
        r += 1

        idxs = compute_factor_indices(model)
        # show in order: higher-order rows (if any) then factor_groups in order
        keys = [k for k in idxs.keys() if k.startswith("HigherOrder:")] + factor_groups
        for keyf in keys:
            vals = idxs.get(keyf, {})
            label = keyf.replace("HigherOrder:", "Higher-order: ")
            _set(ws, r, 1, _factor_display_formula(label) if not label.startswith("Higher-order:") else label, align=ALIGN_LEFT)
            _set(ws, r, 2, vals.get("k"), align=ALIGN_RIGHT, num_format="0")
            _set(ws, r, 3, vals.get("AVE"), align=ALIGN_RIGHT, num_format="0.000")
            _set(ws, r, 4, vals.get("alpha"), align=ALIGN_RIGHT, num_format="0.000")
            _set(ws, r, 5, vals.get("omega"), align=ALIGN_RIGHT, num_format="0.000")
            _set(ws, r, 6, vals.get("omegaH"), align=ALIGN_RIGHT, num_format="0.000")
            _set(ws, r, 7, vals.get("omegaS"), align=ALIGN_RIGHT, num_format="0.000")
            r += 1
        _apply_border_row(ws, r - 1, 1, len(headers), border_bottom())

        # spacing before next model
        r += 2

    # Global note
    note = ("Note. λ = STDYX standardized loading; θ = residual variance (1 - R²). Significant λ values are bolded. "
            "For ESEM/bifactor-ESEM models, all estimated loadings are shown. "
            "Alpha and omega are computed from the model-implied covariance matrix using STDYX loadings, factor correlations, and θ; "
            "if R² is missing for any item, reliability indices are left blank.")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
    _set(ws, r, 1, note, font=FONT_NOTE, align=ALIGN_LEFT)

def build_audit(ws, models: List[ParsedModel]):
    ws.title = "Audit_ParsedValues"
    headers = ["Filename", "Model label", "Type", "chi2", "df", "p", "cfi", "tli", "rmsea", "ci_lo", "ci_hi", "pclose", "srmr", "aic", "bic", "abic", "terminated"]
    for i, h in enumerate(headers, start=1):
        _set(ws, 1, i, h, font=FONT_BOLD, align=ALIGN_CENTER)
    _apply_border_row(ws, 1, 1, len(headers), border_top_bottom())
    _set_col_widths(ws, {1: 40, 2: 18, 3: 22, 4: 10, 5: 6, 6: 10, 7: 8, 8: 8, 9: 8, 10: 8, 11: 8, 12: 10, 13: 8, 14: 12, 15: 12, 16: 12, 17: 12})
    r = 2
    for pm in models:
        fs = pm.fit
        row = [
            pm.out_file.name, pm.model_label, pm.model_type,
            fs.chi2, fs.df, fs.p, fs.cfi, fs.tli, fs.rmsea, fs.rmsea_ci_lo, fs.rmsea_ci_hi, fs.pclose,
            fs.srmr, fs.aic, fs.bic, fs.abic, fs.terminated
        ]
        for c, v in enumerate(row, start=1):
            align = ALIGN_LEFT if c in (1,2,3) else ALIGN_RIGHT
            fmt = None
            if c in (4,6,7,8,9,10,11,12,13,14,15,16):
                fmt = "0.000"
            if c == 5:
                fmt = "0"
            _set(ws, r, c, v, align=align, num_format=fmt)
        r += 1
    if r > 2:
        _apply_border_row(ws, r - 1, 1, len(headers), border_bottom())


def apply_global_format(ws):
    """Apply Times New Roman 10 and consistent alignment to the used range."""
    max_r = ws.max_row or 1
    max_c = ws.max_column or 1
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            cell = ws.cell(r, c)
            if cell.value is None and cell.has_style:
                # keep explicit styles on empty styled cells
                pass
            if cell.font is None or cell.font.name is None:
                cell.font = FONT_BODY
            else:
                # enforce font family/size while keeping bold/italic
                cell.font = Font(
                    name="Times New Roman",
                    size=cell.font.size or 10,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    underline=cell.font.underline,
                    color=cell.font.color
                )
            if cell.alignment is None:
                cell.alignment = ALIGN_CENTER
            # default number format stays as-is

# -----------------------------
# Parse models
# -----------------------------
def parse_models(folder: Path, pattern: str = "*.out") -> List[ParsedModel]:
    manifest = load_manifest(folder)
    out_files = sorted(folder.glob(pattern))

    models: List[ParsedModel] = []
    for fp in out_files:
        txt = _read_text(fp)

        meta = manifest.get(fp.name.lower(), {})
        model_id, inferred_type = _parse_model_id_and_type(fp.name)

        model_label = meta.get("model_label") or (f"Model {model_id}" if model_id is not None else fp.stem)
        model_type = meta.get("model_type") or inferred_type or fp.stem
        role = (meta.get("role") or "").strip().lower()

        order = 10**9
        if meta.get("order"):
            try:
                order = int(meta["order"])
            except Exception:
                order = 10**9
        elif model_id is not None:
            order = model_id

        usevars = set(parse_usevariables(txt))

        all_loadings = parse_stdyx_loadings(txt)
        # split observed vs higher-order (latent indicators)
        observed_loadings: Dict[str, Dict[str, Loading]] = {}
        higher_order: Dict[str, Dict[str, Loading]] = {}
        for fac, inds in all_loadings.items():
            for ind, ld in inds.items():
                if ind in usevars:
                    observed_loadings.setdefault(fac, {})[ind] = ld
                else:
                    # latent indicators (2nd-order)
                    higher_order.setdefault(fac, {})[ind] = ld

        models.append(
            ParsedModel(
                out_file=fp,
                model_label=model_label,
                model_type=model_type,
                fit=parse_fit(txt),
                loadings={f: d for f, d in observed_loadings.items() if d},
                higher_order={f: d for f, d in higher_order.items() if d},
                phi=parse_stdyx_factor_correlations(txt),
                r2=parse_rsquare(txt),
                role=role,
                order=order,
                model_id=model_id,
            )
        )

    models.sort(key=lambda m: (m.order, m.out_file.name.lower()))
    return models

def _collect_factor_names(models: List[ParsedModel]) -> List[str]:
    s: Set[str] = set()
    for m in models:
        for f in m.loadings.keys():
            s.add(f)
    return sorted(s, key=lambda x: x.lower())

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default=".", help="Folder containing Mplus .out files (and optional model_manifest.csv)")
    ap.add_argument("--output", default="Tabulated Results_v7_6.xlsx", help="Output workbook (.xlsx)")
    ap.add_argument("--top", type=int, default=3, help="Number of top models (from Table 1) to include in Table 2")
    ap.add_argument("--pattern", default="*.out", help="Glob pattern for Mplus output files (default: *.out)")
    args = ap.parse_args()

    folder = Path(args.input).resolve()
    setup_logging(folder)
    logging.info("Input folder: %s", folder)
    logging.info("Pattern: %s", args.pattern)
    logging.info("Output workbook: %s", args.output)
    logging.info("Top N models: %s", args.top)
    models = parse_models(folder, pattern=args.pattern)
    logging.info("Parsed %d model(s)", len(models))
    if not models:
        raise FileNotFoundError(f"No .out files found in: {folder}")

    top_models = select_top_models(models, top_n=args.top)
    logging.info("Selected %d top model(s) for Table 2", len(top_models))
    factor_names = _collect_factor_names(top_models)

    wb = Workbook()
    ws0 = wb.active
    build_instructions(ws0, factor_names)

    ws1 = wb.create_sheet()
    build_table1(ws1, models)

    ws2 = wb.create_sheet()
    build_table2(ws2, models, top_models=top_models)

    ws3 = wb.create_sheet()
    build_audit(ws3, models)

    out_path = folder / args.output
    apply_global_format(ws0)
    apply_global_format(ws1)
    apply_global_format(ws2)
    apply_global_format(ws3)

    wb.save(out_path)
    logging.info("Wrote workbook: %s", out_path)
    print(f"Wrote: {out_path}")
    logging.info("Done")
    return 0

if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except FileNotFoundError:
        print("\nERROR: No Mplus .out files were found.")
        print("Place this program in the folder containing your Mplus .out files, or use --input to point to the correct folder.")
        print("A diagnostic log is written to run.log in the input folder (if the folder exists).\n")
        raise
    except PermissionError:
        print("\nERROR: Permission denied while reading outputs or writing the Excel workbook.")
        print("Close the output workbook if it is open, and ensure you have write access to the folder.\n")
        raise
    except Exception:
        print("\nERROR: An unexpected error occurred while generating the Excel tables.")
        print("Please send the run.log file to the developer for troubleshooting.\n")
        raise
